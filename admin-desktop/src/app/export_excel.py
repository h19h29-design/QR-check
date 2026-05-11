from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DEFAULT_ITEM_LABELS = {
    "document": "서류보관상태",
    "cleaning": "청소상태",
    "lighting": "소등상태",
    "fire": "화기단속상태",
    "door": "문단속상태",
}


def export_security_check_xlsx(
    records: list[dict],
    output_path: str | Path,
    title: str = "보안점검표",
    items: list[dict] | None = None,
) -> Path:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "보안점검표"

    export_items = normalize_export_items(records, items)
    headers = [
        "점검일",
        "실명",
        "담당/당직",
        "점검자",
        *[label for _, label in export_items],
        "특이사항",
        "이상여부",
        "관리자확인",
        "관리자확인일시",
        "담당자 서명",
        "관리자 서명",
    ]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1, safe_excel_text(title)).font = Font(size=16, bold=True)
    ws.cell(1, 1).alignment = Alignment(horizontal="center")
    ws.append(headers)

    for record in records:
        status = record.get("status_json", {})
        if isinstance(status, str):
            try:
                status = json.loads(status)
            except json.JSONDecodeError:
                status = {}
        ws.append(
            [
                safe_excel_text(record.get("inspection_date", "")),
                safe_excel_text(record.get("room_name", "")),
                "당직자" if record.get("role_type") == "duty" else "담당자",
                safe_excel_text(record.get("person_name", "")),
                *[safe_excel_text(status.get(key, "이상 무")) for key, _ in export_items],
                safe_excel_text(record.get("remarks", "")),
                "이상 있음" if record.get("abnormal") else "이상 없음",
                "확인" if record.get("admin_verified") else "미확인",
                safe_excel_text(record.get("admin_verified_at", "")),
                "",
                "",
            ]
        )

    header_fill = PatternFill("solid", fgColor="E8EEF7")
    bad_fill = PatternFill("solid", fgColor="FFE5E5")
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    abnormal_index = headers.index("이상여부")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        if row[abnormal_index].value == "이상 있음":
            for cell in row:
                cell.fill = bad_fill
    for cell in ws[2]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [12, 16, 10, 12] + [15] * len(export_items) + [24, 12, 12, 20, 14, 14]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A3"
    wb.save(path)
    return path


def normalize_export_items(records: list[dict], items: list[dict] | None = None) -> list[tuple[str, str]]:
    result: list[tuple[str, str]] = []
    seen: set[str] = set()
    for item in items or []:
        key = str(item.get("item_key") or "").strip()
        label = str(item.get("item_name") or key).strip()
        active = item.get("active", True)
        if key and active not in (False, 0, "0", "false", "FALSE") and key not in seen:
            result.append((key, label or key))
            seen.add(key)
    if not result:
        for key, label in DEFAULT_ITEM_LABELS.items():
            result.append((key, label))
            seen.add(key)
    for record in records:
        status = record.get("status_json", {})
        if isinstance(status, str):
            try:
                status = json.loads(status)
            except json.JSONDecodeError:
                status = {}
        if isinstance(status, dict):
            for key in status:
                key_text = str(key)
                if key_text and key_text not in seen:
                    result.append((key_text, DEFAULT_ITEM_LABELS.get(key_text, key_text)))
                    seen.add(key_text)
    return result


def safe_excel_text(value: object) -> object:
    if not isinstance(value, str):
        return value
    if value and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value
